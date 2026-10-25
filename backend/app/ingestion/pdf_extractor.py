import logging
import os
from pathlib import Path

import fitz
import pytesseract
from docx import Document as DocxDocument
from PIL import Image

logger = logging.getLogger(__name__)

# Optional: set TESSERACT_CMD in your backend .env file if
# Tesseract is not available through your system PATH.
TESSERACT_PATH = os.getenv("TESSERACT_CMD")
if TESSERACT_PATH:
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH


def extract_pdf_pages(file_path: str | Path) -> list[dict]:
    """Extract selectable text and OCR text from PDF pages."""
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError("The uploaded PDF could not be found.")

    pages = []

    try:
        with fitz.open(path) as pdf:
            if pdf.is_encrypted:
                raise ValueError(
                    "Password-protected PDFs are not supported."
                )

            if len(pdf) == 0:
                raise ValueError("The PDF contains no pages.")

            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()

                # Use OCR when the page has no selectable text.
                if not text:
                    pixmap = page.get_pixmap(
                        matrix=fitz.Matrix(2, 2),
                        alpha=False,
                    )

                    image = Image.frombytes(
                        "RGB",
                        (pixmap.width, pixmap.height),
                        pixmap.samples,
                    )

                    try:
                        text = pytesseract.image_to_string(image).strip()
                    except pytesseract.TesseractNotFoundError as exc:
                        raise ValueError(
                            "Tesseract OCR is required to process scanned "
                            "PDF pages. Install Tesseract and configure "
                            "TESSERACT_CMD or your system PATH."
                        ) from exc

                pages.append(
                    {
                        "page_number": page_number,
                        "text": text,
                    }
                )

    except pytesseract.TesseractNotFoundError as exc:
        raise ValueError(
            "Tesseract OCR was not found. Install it and configure "
            "TESSERACT_CMD or your system PATH."
        ) from exc

    except fitz.FileDataError as exc:
        raise ValueError(
            "The PDF is damaged or cannot be read."
        ) from exc

    if not any(page["text"].strip() for page in pages):
        raise ValueError(
            "No text could be extracted from this PDF, even with OCR."
        )

    return pages


def extract_text_pages(file_path: str | Path) -> list[dict]:
    """Extract text from a plain-text file."""
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError("The uploaded text file could not be found.")

    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError(
            "The TXT file must use UTF-8 text encoding."
        ) from exc

    if not text.strip():
        raise ValueError("The TXT file contains no readable text.")

    return [{"page_number": 1, "text": text.strip()}]


def extract_docx_pages(file_path: str | Path) -> list[dict]:
    """Extract paragraphs and table text from a DOCX document."""
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError("The uploaded DOCX file could not be found.")

    try:
        document = DocxDocument(str(path))

        parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if cells:
                    parts.append(" | ".join(cells))

    except Exception as exc:
        raise ValueError(
            "The DOCX file is damaged or cannot be read."
        ) from exc

    text = "\n".join(parts).strip()

    if not text:
        raise ValueError(
            "No readable text was found in the DOCX document."
        )

    return [{"page_number": 1, "text": text}]


def extract_document_pages(file_path: str | Path) -> list[dict]:
    """Extract text from PDF, TXT, or DOCX files."""
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf_pages(path)

    if extension == ".txt":
        return extract_text_pages(path)

    if extension == ".docx":
        return extract_docx_pages(path)

    raise ValueError(
        "Unsupported file type. Upload a PDF, DOCX, or TXT file."
    )